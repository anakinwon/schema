import sys, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
sys.stdout.reconfigure(encoding='utf-8')

EXCEL_PATH = r'c:\Users\anaki\workspace\claude-nextjs-starters\docs\da-plan\테이블정의서.xlsx'

TABLES = [
    {
        'sheet': 'tb_product_category', 'phys': 'tb_product_category', 'logic': '상품분류',
        'cols': [
            (1,'상품분류코드',    'prd_cls_cd',      'VARCHAR',  10,  None, 'N', '',      'PK', '',   '상품 분류 식별코드'),
            (2,'상품분류명',      'prd_cls_nm',      'VARCHAR',  100, None, 'N', '',      '',   '',   '분류 명칭'),
            (3,'상위상품분류코드','upr_prd_cls_cd',  'VARCHAR',  10,  None, 'Y', '',      '',   'FK', '최상위는 NULL'),
            (4,'사용여부',        'use_yn',          'VARCHAR',  1,   None, 'N', 'Y',     '',   '',   'Y사용 N미사용'),
        ]
    },
    {
        'sheet': 'tb_product', 'phys': 'tb_product', 'logic': '상품',
        'cols': [
            (1,'상품번호',         'prd_no',         'VARCHAR',  10,  None, 'N', '',    'PK', '',   '상품 식별번호'),
            (2,'상품명',           'prd_nm',         'VARCHAR',  200, None, 'N', '',    '',   '',   '상품 공식 명칭'),
            (3,'상품분류코드',     'prd_cls_cd',     'VARCHAR',  10,  None, 'N', '',    '',   'FK', '분류 외래키'),
            (4,'상품판매단가',     'prd_sl_up',      'NUMERIC',  21,  3,    'N', '0',   '',   '',   '1단위 판매가'),
            (5,'상품재고수량',     'prd_inv_qty',    'INTEGER',  None,None, 'N', '0',   '',   '',   '현재 재고'),
            (6,'상품상태종류코드', 'prd_sts_knd_cd', 'VARCHAR',  2,   None, 'N', '01',  '',   '',   '01판매중 02품절 03판매중지'),
            (7,'상품내용',         'prd_cont',       'TEXT',     None,None, 'Y', '',    '',   '',   '상품 설명'),
        ]
    },
    {
        'sheet': 'tb_customer', 'phys': 'tb_customer', 'logic': '고객',
        'cols': [
            (1,'고객번호',         'ct_no',          'VARCHAR',  20,  None, 'N', '',    'PK', '',   '고객 식별번호'),
            (2,'고객명',           'ct_nm',          'VARCHAR',  100, None, 'N', '',    '',   '',   '실명/등록명'),
            (3,'고객이메일주소',   'ct_eml_addr',    'VARCHAR',  200, None, 'Y', '',    '',   '',   '이메일(UNIQUE)'),
            (4,'고객비밀번호암호', 'ct_pwd_enc',     'VARCHAR',  500, None, 'N', '',    '',   '',   '암호화 저장'),
            (5,'고객생일',         'ct_bdy',         'VARCHAR',  8,   None, 'Y', '',    '',   '',   'YYYYMMDD'),
            (6,'성별구분코드',     'gnd_gbn_cd',     'VARCHAR',  2,   None, 'Y', '',    '',   '',   '01남 02여 09기타'),
            (7,'고객상태종류코드', 'ct_sts_knd_cd',  'VARCHAR',  2,   None, 'N', '01',  '',   '',   '01정상 02휴면 03탈퇴'),
            (8,'고객기본주소',     'ct_bsic_addr',   'VARCHAR',  200, None, 'Y', '',    '',   '',   '도로명/지번'),
            (9,'고객상세주소',     'ct_dtl_addr',    'VARCHAR',  100, None, 'Y', '',    '',   '',   '동호수 등'),
        ]
    },
    {
        'sheet': 'tb_order', 'phys': 'tb_order', 'logic': '주문',
        'cols': [
            (1,'주문번호',         'od_no',          'VARCHAR',  20,  None, 'N', '',                  'PK', '',   '주문 식별번호'),
            (2,'고객번호',         'ct_no',          'VARCHAR',  20,  None, 'N', '',                  '',   'FK', '고객 외래키'),
            (3,'주문일시',         'od_dts',         'TIMESTAMP',None,None, 'N', 'CURRENT_TIMESTAMP', '',   '',   '주문 접수 일시'),
            (4,'주문합계금액',     'od_tot_amt',     'NUMERIC',  21,  3,    'N', '0',                 '',   '',   '전체 상품 합계'),
            (5,'배송기본주소',     'dlv_bsic_addr',  'VARCHAR',  200, None, 'Y', '',                  '',   '',   '배송지 기본주소'),
            (6,'배송상세주소',     'dlv_dtl_addr',   'VARCHAR',  100, None, 'Y', '',                  '',   '',   '배송지 상세주소'),
            (7,'주문상태종류코드', 'od_sts_knd_cd',  'VARCHAR',  2,   None, 'N', '01',                '',   '',   '01주문접수~07반품'),
        ]
    },
    {
        'sheet': 'tb_order_item', 'phys': 'tb_order_item', 'logic': '주문상품',
        'cols': [
            (1,'주문번호',  'od_no',     'VARCHAR', 20,  None, 'N', '',   'PK', 'FK', '복합PK'),
            (2,'주문순번',  'od_seqno',  'INTEGER', None,None, 'N', '',   'PK', '',   '복합PK 순번'),
            (3,'상품번호',  'prd_no',    'VARCHAR', 10,  None, 'N', '',   '',   'FK', '상품 외래키'),
            (4,'주문수량',  'od_qty',    'INTEGER', None,None, 'N', '1',  '',   '',   '1 이상'),
            (5,'주문단가',  'od_up',     'NUMERIC', 21,  3,    'N', '',   '',   '',   '주문시점 단가'),
            (6,'주문금액',  'od_amt',    'NUMERIC', 21,  3,    'N', '',   '',   '',   'qty*up'),
            (7,'취소여부',  'ccl_yn',    'VARCHAR', 1,   None, 'N', 'N',  '',   '',   'Y취소 N정상'),
        ]
    },
]

SYS_COLS = [
    ('등록사용자아이디','reg_usr_id', 'VARCHAR',  20,  None, 'N', 'ADMIN',              '','','시스템컬럼'),
    ('등록일시',        'reg_dts',   'TIMESTAMP',None,None, 'N', 'CURRENT_TIMESTAMP',  '','','시스템컬럼'),
    ('수정사용자아이디','mod_usr_id', 'VARCHAR',  20,  None, 'N', 'ADMIN',              '','','시스템컬럼'),
    ('수정일시',        'mod_dts',   'TIMESTAMP',None,None, 'N', 'CURRENT_TIMESTAMP',  '','','시스템컬럼'),
]

wb = openpyxl.Workbook()
wb.remove(wb.active)

HEAD_FILL  = PatternFill('solid', fgColor='1F497D')
HEAD_FONT  = Font(bold=True, color='FFFFFF', size=10)
TITLE_FILL = PatternFill('solid', fgColor='2E75B6')
TITLE_FONT = Font(bold=True, color='FFFFFF', size=12)
SYS_FILL   = PatternFill('solid', fgColor='E2EFDA')
BORDER     = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))
CENTER     = Alignment(horizontal='center', vertical='center')
HEADERS    = ['순번','논리명','물리명','데이터타입','길이','소수','NULL여부','DEFAULT','PK','FK','설명']
WIDTHS     = [5, 20, 22, 12, 6, 6, 8, 20, 5, 5, 30]

# 목록 시트
ws0 = wb.create_sheet('테이블목록')
ws0.append(['No','테이블물리명','테이블논리명','비즈니스컬럼수','비고'])
for i, t in enumerate(TABLES, 1):
    ws0.append([i, t['phys'], t['logic'], len(t['cols']), '시스템컬럼 4종 별도'])

# 테이블 시트
for t in TABLES:
    ws = wb.create_sheet(t['sheet'])
    ws.merge_cells('A1:K1')
    ws['A1'] = f"{t['phys']} ({t['logic']})"
    ws['A1'].fill = TITLE_FILL
    ws['A1'].font = TITLE_FONT
    ws['A1'].alignment = CENTER

    ws.append(HEADERS)
    for ci, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=2, column=ci)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = CENTER; cell.border = BORDER
        ws.column_dimensions[chr(64+ci)].width = w

    for row in t['cols']:
        ws.append(list(row))
        for ci in range(1, 12):
            ws.cell(row=ws.max_row, column=ci).border = BORDER

    seq = len(t['cols']) + 1
    for sc in SYS_COLS:
        ws.append([seq] + list(sc))
        for ci in range(1, 12):
            ws.cell(row=ws.max_row, column=ci).fill = SYS_FILL
            ws.cell(row=ws.max_row, column=ci).border = BORDER
        seq += 1

wb.save(EXCEL_PATH)
print(f'[생성완료] 테이블정의서.xlsx')
for sn in wb.sheetnames:
    ws2 = wb[sn]
    print(f'  {sn}: {ws2.max_row}행')
